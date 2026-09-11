import json, re, collections, pathlib, xml.etree.ElementTree as ET
from openpyxl import load_workbook
base = pathlib.Path('D:/系统知识库/00_Inbox/yljhis')
out = pathlib.Path('artifacts/verification/catalog-audit'); out.mkdir(parents=True, exist_ok=True)
rows=[]
for p in (base/'reports').rglob('*'):
    if p.suffix.lower() != '.xml': continue
    b=p.read_bytes(); head=b[:150].decode('ascii','ignore'); m=re.search(r'encoding=[\"\']([^\"\']+)',head)
    enc=m[1] if m else ('utf-16' if b.startswith((b'\xff\xfe',b'\xfe\xff')) else 'utf-8-sig')
    try: s=b.decode(enc)
    except (UnicodeError, LookupError): s=b.decode('gb18030',errors='replace')
    rootmatch=re.search(r'<(?![?!])([\w:]+)',s); rootname=rootmatch[1] if rootmatch else '?'
    row={'file':p.name,'root':rootname}
    if rootname=='ReportQueryInfo':
        root=ET.fromstring(s); sources=root.findall('./QueryDataSource/QueryDataSource')
        row['sources']=[{'name':n.findtext('Name'),'kind':n.findtext('SqlType'),'hasSql':bool((n.findtext('Sql') or '').strip())} for n in sources]
        row['groupSql']=bool((root.findtext('./TableGroup/QueryDataSource/Sql') or '').strip())
    rows.append(row)
(out/'xml-structure.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf-8')
print('ROOTS',collections.Counter(r['root'] for r in rows))
empty=[r for r in rows if r['root']=='ReportQueryInfo' and not any(n['hasSql'] for n in r['sources'])]
print('EMPTY',json.dumps(empty,ensure_ascii=False))
print('OTHER',[(r['file'],r['root']) for r in rows if r['root'] not in ('ReportQueryInfo','Spread')][:35])
for p in (base/'docs').glob('*.xlsx'):
    wb=load_workbook(p,read_only=True,data_only=True)
    print('WORKBOOK',p.name,'SHEETS',wb.sheetnames)
    hits=[]
    for ws in wb:
        for i,row in enumerate(ws.iter_rows(values_only=True),1):
            vals=[str(v) for v in row if v is not None]
            if any(re.search(r'菜单|menu|报表中心|各职能科室|抗菌药物查询',v,re.I) for v in vals):
                hits.append({'sheet':ws.title,'row':i,'values':vals})
    print('MENU_HITS',json.dumps(hits[:30],ensure_ascii=False))
    (out/(p.stem+'-menu-hits.json')).write_text(json.dumps(hits,ensure_ascii=False,indent=2),encoding='utf-8')
    wb.close()
